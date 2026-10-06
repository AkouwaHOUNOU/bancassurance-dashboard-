"""Parcours réels AppTest : chaque test travaille sur une copie SQLite isolée."""
import hashlib
import os
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from historical_service import HistoricalService
from streamlit.testing.v1 import AppTest


class SimplifiedAppTests(unittest.TestCase):
    def setUp(self):
        self.production_hash = hashlib.sha256((ROOT / 'historical_bancassurance.db').read_bytes()).hexdigest()
        scratch = Path(os.environ.get('TMPDIR', Path.home() / 'AppData/Local/hermes/cache/scratch'))
        scratch.mkdir(parents=True, exist_ok=True)
        self.temp = tempfile.TemporaryDirectory(dir=scratch)
        self.db = Path(self.temp.name) / 'test.db'
        # sqlite3.__exit__ ne ferme pas la connexion, ce qui verrouille le fichier sous Windows.
        source = sqlite3.connect(ROOT / 'historical_bancassurance.db')
        target = sqlite3.connect(self.db)
        try:
            source.backup(target)
        finally:
            target.close()
            source.close()
        self.env = patch.dict(os.environ, {'BANCASSURANCE_DB_PATH': str(self.db)})
        self.env.start()
        self.service = HistoricalService(str(self.db))

    def tearDown(self):
        self.env.stop()
        self.assertEqual(self.production_hash, hashlib.sha256((ROOT / 'historical_bancassurance.db').read_bytes()).hexdigest())
        self.temp.cleanup()

    def navigate(self, app, page):
        app.sidebar.radio(key="navigation").set_value(page).run()
        self.assertFalse(app.exception, str(app.exception))
        return app

    def app(self):
        app = AppTest.from_file(str(ROOT / 'app.py'), default_timeout=30).run()
        self.assertFalse(app.exception, str(app.exception))
        return app

    def test_comparison_chart_calendar_order(self):
        from database_service import MOIS
        self.service.save_values_batch('AFRICAN LEASE', 2021, {
            'Janvier': {'objectif': 100, 'realisation': 50},
            'Février': {'objectif': 100, 'realisation': 70},
        })
        with patch('streamlit.altair_chart') as charts:
            app = self.navigate(self.app(), 'Comparaison')
            charts.reset_mock()
            app.sidebar.number_input(key='global_year').set_value(2021).run()
            specs = [call.args[0].to_dict() for call in charts.call_args_list]
        lines = [spec for spec in specs if spec.get('mark', {}).get('type') == 'line']
        self.assertTrue(lines, 'La comparaison doit utiliser un graphique chronologique explicite')
        self.assertEqual(lines[0]['encoding']['x']['sort'], MOIS)

    def test_month_save_reload_twelve_months_preserves_manual_annual(self):
        self.service.remove_annee('AFRICAN LEASE', 2021)
        self.service.save_values_batch('AFRICAN LEASE', 2021, {'ANNUEL': {'objectif': 9000, 'realisation': 7000}})
        app = self.navigate(self.app(), 'Production')
        app.sidebar.number_input(key='global_year').set_value(2021).run()
        self.assertIn('Mois à saisir', [widget.label for widget in app.selectbox])
        self.assertEqual(app.selectbox(key='month_AFRICAN LEASE_2021').options, ['Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin', 'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre'])
        app.number_input(key='monthly_obj_AFRICAN LEASE_2021_Janvier').set_value(1234)
        app.number_input(key='monthly_real_AFRICAN LEASE_2021_Janvier').set_value(789)
        app.button(key='save_month_AFRICAN LEASE_2021_Janvier').click().run()
        self.assertFalse(app.exception)
        values = self.service.get_data_for_partner_annee('AFRICAN LEASE', 2021)
        self.assertEqual(values['Janvier'], {'objectif': 1234.0, 'realisation': 789.0})
        self.assertEqual(values['ANNUEL'], {'objectif': 9000.0, 'realisation': 7000.0})
        reload = self.navigate(self.app(), 'Production')
        reload.sidebar.number_input(key='global_year').set_value(2021).run()
        self.assertEqual(reload.number_input(key='monthly_obj_AFRICAN LEASE_2021_Janvier').value, 1234)
        reload.selectbox(key='month_AFRICAN LEASE_2021').select('Décembre').run()
        self.assertIsNone(reload.number_input(key='monthly_obj_AFRICAN LEASE_2021_Décembre').value)
        reload.sidebar.selectbox(key='global_partner').select('BIA TOGO').run()
        self.assertFalse(reload.exception)

    def test_annual_manual_and_explicit_monthly_calculation(self):
        self.service.remove_annee('AFRICAN LEASE', 2021)
        self.service.save_values_batch('AFRICAN LEASE', 2021, {'Janvier': {'objectif': 1200, 'realisation': 800}})
        app = self.navigate(self.app(), 'Production')
        app.sidebar.number_input(key='global_year').set_value(2021).run()
        self.assertIn('Mode du total annuel', [w.label for w in app.radio])
        app.radio(key='annual_mode_AFRICAN LEASE_2021').set_value('Totaux annuels manuels').run()
        app.number_input(key='annual_obj_AFRICAN LEASE_2021').set_value(5000)
        app.number_input(key='annual_real_AFRICAN LEASE_2021').set_value(0)
        app.button(key='save_annual_AFRICAN LEASE_2021').click().run()
        self.navigate(app, 'Tableau de bord')
        self.assertEqual([m.value for m in app.metric][:3], ['5,000 FCFA', '0 FCFA', '5,000 FCFA'])
        reload = self.navigate(self.app(), 'Production')
        reload.sidebar.number_input(key='global_year').set_value(2021).run()
        self.assertEqual(reload.radio(key='annual_mode_AFRICAN LEASE_2021').value, 'Totaux annuels manuels')
        reload.radio(key='annual_mode_AFRICAN LEASE_2021').set_value('Somme des mois').run()
        reload.button(key='save_annual_AFRICAN LEASE_2021').click().run()
        self.assertEqual(self.service.get_gap_taux_historical('AFRICAN LEASE', 2021)['objectif_N'], 1200)
        self.assertEqual(self.service.get_profil_partenaire('AFRICAN LEASE', 2021)['realisation_annuelle'], 800)

    def test_weekly_first_entry_save_reload_on_empty_year(self):
        self.service.remove_annee('AFRICAN LEASE', 2021)
        app = self.navigate(self.app(), 'Suivi hebdomadaire')
        app.sidebar.number_input(key='global_year').set_value(2021).run()
        self.assertIn('Mois hebdomadaire', [w.label for w in app.selectbox])
        app.selectbox(key='weekly_month_AFRICAN LEASE_2021').select('Février').run()
        app.selectbox(key='weekly_week_AFRICAN LEASE_2021_Février').select(5).run()
        app.number_input(key='weekly_obj_AFRICAN LEASE_2021_Février_5').set_value(222)
        app.number_input(key='weekly_real_AFRICAN LEASE_2021_Février_5').set_value(111)
        app.text_input(key='weekly_comment_AFRICAN LEASE_2021_Février_5').set_value('Visite réalisée')
        app.button(key='save_week_AFRICAN LEASE_2021_Février_5').click().run()
        self.assertFalse(app.exception)
        self.assertEqual(self.service.get_weekly_for_partner_mois('AFRICAN LEASE', 2021, 'Février')[5], {'objectif': 222.0, 'realisation': 111.0, 'commentaire': 'Visite réalisée'})
        reload = self.navigate(self.app(), 'Suivi hebdomadaire')
        reload.sidebar.number_input(key='global_year').set_value(2021).run()
        reload.selectbox(key='weekly_month_AFRICAN LEASE_2021').select('Février').run()
        reload.selectbox(key='weekly_week_AFRICAN LEASE_2021_Février').select(5).run()
        self.assertEqual(reload.text_input(key='weekly_comment_AFRICAN LEASE_2021_Février_5').value, 'Visite réalisée')

    def test_commissions_use_dedicated_text_fields_save_reload(self):
        self.service.save_values_batch('AFRICAN LEASE', 2021, {'ANNUEL': {'objectif': 9000, 'realisation': 7000}})
        self.service.save_text('AFRICAN LEASE', 2021, 'commission_reversee', '125 000')
        app = self.navigate(self.app(), 'Commissions')
        app.sidebar.number_input(key='global_year').set_value(2021).run()
        self.assertIn('Commission reversée (FCFA)', [w.label for w in app.text_input])
        self.assertEqual(app.text_input(key='commission_reversee_AFRICAN LEASE_2021').value, '125 000')
        app.text_input(key='commission_reversee_AFRICAN LEASE_2021').set_value('250 000')
        app.text_input(key='total_a_payer_AFRICAN LEASE_2021').set_value('37 500')
        app.text_input(key='date_traitement_AFRICAN LEASE_2021').set_value('2021-02-03')
        app.text_input(key='date_paiement_AFRICAN LEASE_2021').set_value('2021-02-05')
        app.button(key='save_commissions_AFRICAN LEASE_2021').click().run()
        self.assertEqual(self.service.get_data_for_partner_annee('AFRICAN LEASE', 2021)['ANNUEL'], {'objectif': 9000.0, 'realisation': 7000.0})
        reload = self.navigate(self.app(), 'Commissions')
        reload.sidebar.number_input(key='global_year').set_value(2021).run()
        self.assertEqual(reload.text_input(key='total_a_payer_AFRICAN LEASE_2021').value, '37 500')
        self.assertEqual(reload.text_input(key='date_paiement_AFRICAN LEASE_2021').value, '2021-02-05')
        self.assertFalse(reload.exception)

    def test_profile_single_form_all_fields_save_reload(self):
        fields = ['specifications', 'responsables_bancaires', 'dg', 'points_forts', 'points_faibles', 'points_focaux', 'personnes_ressources', 'activites_planifiees', 'cibles', 'actions_trimestre', 'a_fin_septembre', 'difficultes', 'ameliorations', 'recommandations', 'observations']
        app = self.navigate(self.app(), 'Profil')
        app.sidebar.number_input(key='global_year').set_value(2021).run()
        self.assertIn('DG du partenaire', [w.label for w in app.text_area])
        for field in fields:
            app.text_area(key=f'profile_{field}_AFRICAN LEASE_2021').set_value('Test ' + field)
        app.button(key='save_profile_AFRICAN LEASE_2021').click().run()
        self.assertFalse(app.exception)
        stored = self.service.get_text_for_partner_annee('AFRICAN LEASE', 2021)
        for field in fields:
            self.assertEqual(stored[field], 'Test ' + field)
        reload = self.navigate(self.app(), 'Profil')
        reload.sidebar.number_input(key='global_year').set_value(2021).run()
        for field in fields:
            self.assertEqual(reload.text_area(key=f'profile_{field}_AFRICAN LEASE_2021').value, 'Test ' + field)
        reload.sidebar.number_input(key='global_year').set_value(2022).run()
        self.assertNotEqual(reload.text_area(key='profile_dg_AFRICAN LEASE_2022').value, 'Test dg')
        self.assertEqual(len([w for w in reload.text_area if w.label == 'DG du partenaire']), 1)

    def test_dynamic_partner_persists_without_fake_values(self):
        self.service.save_text('PARTENAIRE EXISTANT', 2020, 'dg', 'Contact')
        app = self.app()
        self.assertIn('PARTENAIRE EXISTANT', app.sidebar.selectbox(key='global_partner').options)
        app.sidebar.text_input(key='new_partner_name').set_value('NOUVEAU PARTENAIRE')
        app.sidebar.button(key='add_partner').click().run()
        self.assertFalse(app.exception)
        reload = self.app()
        self.assertIn('NOUVEAU PARTENAIRE', reload.sidebar.selectbox(key='global_partner').options)
        reload.sidebar.selectbox(key='global_partner').select('NOUVEAU PARTENAIRE').run()
        self.assertFalse(reload.exception)
        self.assertEqual(self.service.get_data_for_partner_annee('NOUVEAU PARTENAIRE', 2021), {})
        self.assertEqual(self.service.get_all_annees('PARTENAIRE EXISTANT'), [2020])
        self.service.save_weekly('NOUVEAU PARTENAIRE', 2019, 'Janvier', 1, 0, 0, '')
        self.assertEqual(self.service.get_all_annees('NOUVEAU PARTENAIRE'), [2019])

    def test_comparison_read_only_global_year_and_effective_totals(self):
        self.service.remove_annee('AFRICAN LEASE', 2021)
        self.service.save_values_batch('AFRICAN LEASE', 2021, {'Janvier': {'objectif': 1200, 'realisation': 800}})
        self.service.save_weekly('AFRICAN LEASE', 2021, 'Février', 5, 222, 111, 'Visite')
        app = self.navigate(self.app(), 'Comparaison')
        app.sidebar.number_input(key='global_year').set_value(2021).run()
        self.assertIn('Années à comparer', [w.label for w in app.multiselect])
        self.assertIn(2021, app.multiselect(key='compare_years_AFRICAN LEASE_2021').value)
        self.assertEqual(len(app.number_input), 1)
        self.assertEqual(len(app.text_area), 0)
        comparison = self.service.get_comparison_dataframe('AFRICAN LEASE', [2021])
        self.assertEqual(comparison.loc[comparison.Mois == 'ANNUEL', 'Objectif 2021'].iloc[0], 1200)
        weekly = self.service.get_weekly_comparison_dataframe('AFRICAN LEASE', [2021], 'Février')
        self.assertEqual(weekly.loc[weekly.Semaine == 'S5', 'Réalisation 2021'].iloc[0], 111)
        app.multiselect(key='compare_years_AFRICAN LEASE_2021').set_value([]).run()
        self.assertFalse(app.exception)

    def test_export_excel_and_pptx_actual_content_selected_year(self):
        from io import BytesIO
        import importlib.util
        self.assertIsNotNone(importlib.util.find_spec('reporting'), 'Module export testable manquant')
        from reporting import export_excel, export_presentation
        import openpyxl
        from pptx import Presentation
        self.service.remove_annee('AFRICAN LEASE', 2021)
        self.service.save_values_batch('AFRICAN LEASE', 2021, {'Janvier': {'objectif': 1234, 'realisation': 789}, 'ANNUEL': {'objectif': 9000, 'realisation': 7000}})
        self.service.save_text('AFRICAN LEASE', 2021, 'commission_reversee', '125 000')
        self.service.save_text('AFRICAN LEASE', 2021, 'dg', 'DG EXERCICE 2021')
        self.service.save_weekly('AFRICAN LEASE', 2021, 'Février', 5, 222, 111, 'Visite')
        workbook = openpyxl.load_workbook(BytesIO(export_excel(self.service, 'AFRICAN LEASE', [2021])), data_only=True)
        rows = list(workbook['Historique'].values)
        self.assertEqual(rows[0], ('Mois', 'Objectif 2021', 'Réalisation 2021', 'Taux 2021'))
        self.assertEqual(rows[1][1:3], (1234, 789))
        self.assertEqual(rows[-1][1:3], (9000, 7000))
        deck = Presentation(BytesIO(export_presentation(self.service, 'AFRICAN LEASE', 2021)))
        all_text = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn('2021', all_text)
        self.assertIn('9,000 FCFA', all_text)
        self.assertIn('125 000', all_text)
        self.assertIn('DG EXERCICE 2021', all_text)
        charts = [shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]
        self.assertTrue(charts)
        self.assertEqual(list(charts[0].series[0].values)[0], 1234)
        self.service.save_text('AFRICAN LEASE', 2021, 'annual_mode', 'monthly')
        workbook_auto = openpyxl.load_workbook(BytesIO(export_excel(self.service, 'AFRICAN LEASE', [2021])), data_only=True)
        self.assertEqual(list(workbook_auto['Historique'].values)[-1][1:3], (1234, 789))
        app = self.navigate(self.app(), 'Comparaison')
        app.sidebar.number_input(key='global_year').set_value(2021).run()
        app.button(key='export_excel_AFRICAN LEASE_2021').click().run()
        app.button(key='export_pptx_AFRICAN LEASE_2021').click().run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.get('download_button')), 2)

    def test_delete_year_requires_exact_confirmation_is_scoped(self):
        self.service.save_values_batch('AFRICAN LEASE', 2021, {'Janvier': {'objectif': 1234, 'realisation': 789}})
        self.service.save_text('AFRICAN LEASE', 2021, 'dg', 'À supprimer')
        self.service.save_weekly('AFRICAN LEASE', 2021, 'Janvier', 1, 12, 7, '')
        self.service.save_value('AFRICAN LEASE', 2022, 'Janvier', 'objectif', 4567)
        app = self.navigate(self.app(), 'Profil')
        app.sidebar.number_input(key='global_year').set_value(2021).run()
        self.assertIn('Confirmation de suppression', [w.label for w in app.text_input])
        self.assertTrue(app.button(key='delete_year_AFRICAN LEASE_2021').disabled)
        app.text_input(key='delete_confirmation_AFRICAN LEASE_2021').set_value('2021').run()
        self.assertTrue(app.button(key='delete_year_AFRICAN LEASE_2021').disabled)
        app.text_input(key='delete_confirmation_AFRICAN LEASE_2021').set_value('AFRICAN LEASE / 2021').run()
        self.assertFalse(app.button(key='delete_year_AFRICAN LEASE_2021').disabled)
        app.button(key='delete_year_AFRICAN LEASE_2021').click().run()
        self.assertFalse(app.exception)
        self.assertEqual(self.service.get_data_for_partner_annee('AFRICAN LEASE', 2021), {})
        self.assertEqual(self.service.get_text_for_partner_annee('AFRICAN LEASE', 2021), {})
        self.assertEqual(self.service.get_weekly_for_partner_mois('AFRICAN LEASE', 2021, 'Janvier'), {})
        self.assertEqual(self.service.get_data_for_partner_annee('AFRICAN LEASE', 2022)['Janvier']['objectif'], 4567)
        self.assertTrue(app.success)
        self.navigate(app, 'Production')
        self.assertIsNone(app.number_input(key='monthly_obj_AFRICAN LEASE_2021_Janvier').value)

    def test_text_form_save_is_atomic_on_sqlite_failure(self):
        self.assertTrue(callable(getattr(self.service, 'save_text_batch', None)))
        self.service.save_text('AFRICAN LEASE', 2021, 'dg', 'Avant')
        with self.assertRaises(sqlite3.ProgrammingError):
            self.service.save_text_batch('AFRICAN LEASE', 2021, {'dg': 'Après', 'cibles': object()})
        self.assertEqual(self.service.get_text_for_partner_annee('AFRICAN LEASE', 2021)['dg'], 'Avant')
        self.service.save_text_batch('AFRICAN LEASE', 2021, {'dg': 'Après', 'cibles': 'Visites'})
        self.assertEqual(self.service.get_text_for_partner_annee('AFRICAN LEASE', 2021)['cibles'], 'Visites')

    def test_existing_partners_all_contexts_read_only(self):
        app = self.app()
        with sqlite3.connect(self.db) as connection:
            before = {table: connection.execute(f'SELECT * FROM {table} ORDER BY 1,2,3').fetchall() for table in ['historical_data', 'textual_data', 'weekly_data']}
        connection.close()
        for partner in self.service.get_all_partenaires():
            app.sidebar.selectbox(key='global_partner').select(partner).run()
            for year in sorted(set(self.service.get_all_annees(partner) + [2021])):
                app.sidebar.number_input(key='global_year').set_value(year).run()
                self.assertFalse(app.exception, f'{partner} / {year}: {app.exception}')
                self.assertEqual(len(app.metric), 4)
        with sqlite3.connect(self.db) as connection:
            after = {table: connection.execute(f'SELECT * FROM {table} ORDER BY 1,2,3').fetchall() for table in before}
        self.assertEqual(before, after)
        # Fermeture explicite nécessaire même après le bloc « with » SQLite.
        connection.close()

    def test_prepared_exports_invalidated_after_save(self):
        app = self.navigate(self.app(), 'Comparaison')
        app.sidebar.number_input(key='global_year').set_value(2021).run()
        app.button(key='export_excel_AFRICAN LEASE_2021').click().run()
        self.assertEqual(len(app.get('download_button')), 1)
        self.navigate(app, 'Production')
        app.number_input(key='monthly_obj_AFRICAN LEASE_2021_Janvier').set_value(9876)
        app.button(key='save_month_AFRICAN LEASE_2021_Janvier').click().run()
        self.assertTrue(app.success)
        self.navigate(app, 'Comparaison')
        self.assertEqual(len(app.get('download_button')), 0)

    def test_empty_comparison_is_clear_without_none_table_or_empty_chart(self):
        app = self.app()
        app.sidebar.number_input(key='global_year').set_value(2040).run()
        self.assertEqual([metric.value for metric in app.metric], ['—'] * 4)
        self.navigate(app, 'Comparaison')
        self.assertEqual(len(app.dataframe), 0)
        self.assertTrue(any('Aucune donnée chiffrée' in info.value for info in app.info))
        self.assertEqual(len(app.get('arrow_vega_lite_chart')), 0)
        app.sidebar.number_input(key='global_year').set_value(2026).run()
        for frame in app.dataframe:
            self.assertFalse(frame.value.isna().any().any())

    def test_light_theme_readable_typography_tokens(self):
        import tomllib
        config = tomllib.loads((ROOT / '.streamlit/config.toml').read_text(encoding='utf-8'))
        self.assertEqual(config['theme'].get('base'), 'light')
        self.assertEqual(config['theme']['backgroundColor'], '#F4F6F8')
        self.assertEqual(config['theme']['textColor'], '#002B49')
        self.assertEqual(config['theme']['primaryColor'], '#D4A72C')
        css = (ROOT / 'styles.css').read_text(encoding='utf-8')
        for token in ['--nsia-page: #F4F6F8', '--nsia-surface: #FFFFFF', '--nsia-text: #002B49', '--nsia-accent: #D4A72C', '--nsia-sidebar: #002B49', 'font-size:18px', 'font-size:16px', 'line-height:1.6', 'width:310px', 'white-space:normal']:
            self.assertIn(token, css)

    def test_navigation_six_pages_global_context_no_sidebar_kpi(self):
        app = self.app()
        self.assertEqual(len(app.tabs), 0)
        self.assertEqual(app.sidebar.radio(key='navigation').options, ['Tableau de bord', 'Production', 'Suivi hebdomadaire', 'Commissions', 'Profil', 'Comparaison'])
        self.assertFalse(any(w.label == 'Mois à saisir' for w in app.selectbox))
        for page, label in [('Production', 'Mois à saisir'), ('Suivi hebdomadaire', 'Mois hebdomadaire'), ('Commissions', 'Commission reversée (FCFA)'), ('Profil', 'DG du partenaire'), ('Comparaison', 'Années à comparer')]:
            app.sidebar.radio(key='navigation').set_value(page).run()
            self.assertFalse(app.exception)
            self.assertEqual(len(app.metric), 0)
            widgets = list(app.selectbox) + list(app.text_input) + list(app.text_area) + list(app.multiselect)
            self.assertIn(label, [w.label for w in widgets])
        app.sidebar.radio(key='navigation').set_value('Tableau de bord').run()
        self.assertEqual(len(app.sidebar.metric), 0)
        self.assertEqual(len(app.metric), 4)
        self.assertEqual(app.sidebar.selectbox(key='global_partner').label, 'Partenaire')
        self.assertEqual(app.sidebar.number_input(key='global_year').label, 'Année')
        app.sidebar.number_input(key='global_year').set_value(2021).run()
        app.sidebar.selectbox(key='global_partner').select('BIA TOGO').run()
        self.assertFalse(app.exception)
        self.assertIn('2021', app.title[0].value)
        self.assertNotIn('Sauvegarder saisies (session)', [button.label for button in app.button])
        self.assertFalse(any('zéro toutes' in button.label for button in app.button))


if __name__ == '__main__':
    unittest.main(verbosity=2)
