"""Suppressions réelles sur copie SQLite ; hash production vérifié à chaque test."""
import unittest
import sqlite3
from test_simplified_app import SimplifiedAppTests as _Base

class AmountDeletionTests(unittest.TestCase):
    setUp = _Base.setUp
    tearDown = _Base.tearDown
    app = _Base.app
    navigate = _Base.navigate

    def test_numeric_deletion_scope_and_annual_fallback(self):
        s = self.service
        for partner, year in [('AFRICAN LEASE', 2021), ('BIA TOGO', 2021), ('AFRICAN LEASE', 2022)]:
            s.save_values_batch(partner, year, {'Janvier': {'objectif': 100, 'realisation': 0}, 'Février': {'objectif': 20}, 'ANNUEL': {'objectif': 900, 'realisation': 700}})
        s.save_text('AFRICAN LEASE', 2021, 'annual_mode', 'manual')
        self.assertTrue(callable(getattr(s, 'delete_amounts', None)), 'Suppression ciblée manquante')
        s.delete_amounts('AFRICAN LEASE', 2021, 'Janvier', ['objectif'])
        self.assertEqual(s.get_data_for_partner_annee('AFRICAN LEASE', 2021)['Janvier'], {'objectif': None, 'realisation': 0})
        s.delete_amounts('AFRICAN LEASE', 2021, 'ANNUEL', ['objectif'])
        self.assertEqual(s.get_gap_taux_historical('AFRICAN LEASE', 2021)['objectif_N'], 20)
        self.assertEqual(s.get_data_for_partner_annee('AFRICAN LEASE', 2021)['ANNUEL']['realisation'], 700)
        self.assertEqual(s.get_text_for_partner_annee('AFRICAN LEASE', 2021)['annual_mode'], 'manual')
        for partner, year in [('BIA TOGO', 2021), ('AFRICAN LEASE', 2022)]:
            self.assertEqual(s.get_data_for_partner_annee(partner, year)['Janvier']['objectif'], 100)
        with self.assertRaises(ValueError):
            s.delete_amounts('AFRICAN LEASE', 2021, 'ANNUEL', ['realisation', 'dg'])
        self.assertEqual(s.get_data_for_partner_annee('AFRICAN LEASE', 2021)['ANNUEL']['realisation'], 700)

    def test_weekly_nullable_deletion_preserves_comment_and_zero(self):
        s = self.service
        s.save_weekly('AFRICAN LEASE', 2021, 'Janvier', 1, 120, 0, 'À garder')
        s.save_weekly('AFRICAN LEASE', 2021, 'Janvier', 2, 33, 22, '')
        self.assertTrue(callable(getattr(s, 'delete_weekly_amounts', None)), 'Suppression hebdomadaire manquante')
        s.delete_weekly_amounts('AFRICAN LEASE', 2021, 'Janvier', 1, ['objectif'])
        self.assertEqual(s.get_weekly_for_partner_mois('AFRICAN LEASE', 2021, 'Janvier')[1], {'objectif': None, 'realisation': 0, 'commentaire': 'À garder'})
        s.delete_weekly_amounts('AFRICAN LEASE', 2021, 'Janvier', 1, ['realisation'])
        s.save_weekly('AFRICAN LEASE', 2021, 'Janvier', 1, None, None, 'Commentaire modifié')
        self.assertIsNone(s.get_weekly_for_partner_mois('AFRICAN LEASE', 2021, 'Janvier')[1]['objectif'])
        s.delete_weekly_amounts('AFRICAN LEASE', 2021, 'Janvier', 2, ['objectif', 'realisation'])
        self.assertNotIn(2, s.get_weekly_for_partner_mois('AFRICAN LEASE', 2021, 'Janvier'))
        with self.assertRaises(ValueError):
            s.delete_weekly_amounts('AFRICAN LEASE', 2021, 'Janvier', 1, ['commentaire'])

    def test_commission_deletion_preserves_dates_and_scope(self):
        s = self.service
        s.save_text_batch('AFRICAN LEASE', 2021, {'commission_reversee': '0', 'total_a_payer': '123 456', 'date_traitement': '2021-01-01', 'date_paiement': '2021-02-01'})
        self.assertTrue(callable(getattr(s, 'delete_commission_amounts', None)), 'Suppression commission manquante')
        s.delete_commission_amounts('AFRICAN LEASE', 2021, ['total_a_payer'])
        values = s.get_text_for_partner_annee('AFRICAN LEASE', 2021)
        self.assertNotIn('total_a_payer', values)
        self.assertEqual(values['commission_reversee'], '0')
        self.assertEqual(values['date_traitement'], '2021-01-01')
        with self.assertRaises(ValueError):
            s.delete_commission_amounts('AFRICAN LEASE', 2021, ['commission_reversee', 'date_paiement'])
        self.assertEqual(s.get_text_for_partner_annee('AFRICAN LEASE', 2021), values)

    def test_both_metrics_scoped_and_atomic_with_sqlite_trigger(self):
        s = self.service
        for partner, year, month, week in [('AFRICAN LEASE', 2021, 'Janvier', 1), ('BIA TOGO', 2021, 'Janvier', 1), ('AFRICAN LEASE', 2022, 'Janvier', 1), ('AFRICAN LEASE', 2021, 'Février', 1), ('AFRICAN LEASE', 2021, 'Janvier', 2)]:
            s.save_weekly(partner, year, month, week, 0, 55, '')
        s.save_text_batch('BIA TOGO', 2021, {'commission_reversee': '55', 'total_a_payer': '66'})
        s.save_text_batch('AFRICAN LEASE', 2022, {'commission_reversee': '55', 'total_a_payer': '66'})
        s.save_text_batch('AFRICAN LEASE', 2021, {'commission_reversee': '55', 'total_a_payer': '66'})
        s.delete_weekly_amounts('AFRICAN LEASE', 2021, 'Janvier', 1, ['objectif', 'realisation'])
        self.assertNotIn(1, s.get_weekly_for_partner_mois('AFRICAN LEASE', 2021, 'Janvier'))
        for partner, year, month, week in [('BIA TOGO', 2021, 'Janvier', 1), ('AFRICAN LEASE', 2022, 'Janvier', 1), ('AFRICAN LEASE', 2021, 'Février', 1), ('AFRICAN LEASE', 2021, 'Janvier', 2)]:
            self.assertEqual(s.get_weekly_for_partner_mois(partner, year, month)[week]['objectif'], 0)
        s.delete_commission_amounts('AFRICAN LEASE', 2021, ['commission_reversee', 'total_a_payer'])
        for partner, year in [('BIA TOGO', 2021), ('AFRICAN LEASE', 2022)]:
            self.assertEqual(s.get_text_for_partner_annee(partner, year)['total_a_payer'], '66')
        self.assertNotIn('total_a_payer', s.get_text_for_partner_annee('AFRICAN LEASE', 2021))
        s.save_values_batch('AFRICAN LEASE', 2021, {'Mars': {'objectif': 0, 'realisation': 10}})
        c = sqlite3.connect(self.db)
        try:
            c.execute("CREATE TRIGGER fail_second_delete BEFORE DELETE ON historical_data WHEN OLD.type='realisation' BEGIN SELECT RAISE(ABORT, 'test rollback'); END")
            c.commit()
        finally:
            c.close()
        with self.assertRaises(sqlite3.IntegrityError):
            s.delete_amounts('AFRICAN LEASE', 2021, 'Mars', ['objectif', 'realisation'])
        self.assertEqual(s.get_data_for_partner_annee('AFRICAN LEASE', 2021)['Mars'], {'objectif': 0, 'realisation': 10})

    def test_ui_deletion_confirmation_reload_and_exports(self):
        s = self.service
        s.remove_annee('AFRICAN LEASE', 2021)
        s.save_values_batch('AFRICAN LEASE', 2021, {'Janvier': {'objectif': 100, 'realisation': 0}, 'ANNUEL': {'objectif': 900, 'realisation': 700}})
        s.save_text_batch('AFRICAN LEASE', 2021, {'annual_mode': 'manual', 'commission_reversee': '125 000', 'total_a_payer': '0', 'date_paiement': '2021-02-01'})
        s.save_weekly('AFRICAN LEASE', 2021, 'Janvier', 1, 12, 0, 'Conserver')
        app = self.navigate(self.app(), 'Comparaison')
        app.sidebar.number_input(key='global_year').set_value(2021).run()
        app.button(key='export_excel_AFRICAN LEASE_2021').click().run()
        for page, suffix, field, input_key in [
            ('Production', 'monthly_AFRICAN LEASE_2021_Janvier', 'Objectif', 'monthly_obj_AFRICAN LEASE_2021_Janvier'),
            ('Production', 'annual_AFRICAN LEASE_2021', 'Objectif', 'annual_obj_AFRICAN LEASE_2021'),
            ('Suivi hebdomadaire', 'weekly_AFRICAN LEASE_2021_Janvier_1', 'Objectif', 'weekly_obj_AFRICAN LEASE_2021_Janvier_1'),
            ('Commissions', 'commissions_AFRICAN LEASE_2021', 'Commission reversée', 'commission_reversee_AFRICAN LEASE_2021')]:
            self.navigate(app, page)
            self.assertTrue(any(w.key == 'delete_amount_' + suffix for w in app.button), 'Bouton suppression ciblée absent')
            self.assertTrue(app.button(key='delete_amount_' + suffix).disabled, suffix + ' ' + str(s.get_data_for_partner_annee('AFRICAN LEASE', 2021)))
            app.selectbox(key='delete_metric_' + suffix).select(field).run()
            app.checkbox(key='confirm_amount_' + suffix + '_' + field).check().run()
            self.assertFalse(app.button(key='delete_amount_' + suffix).disabled)
            app.button(key='delete_amount_' + suffix).click().run()
            self.assertFalse(app.exception, str(app.exception))
            self.assertTrue(app.button(key='delete_amount_' + suffix).disabled, suffix + ' ' + str(s.get_data_for_partner_annee('AFRICAN LEASE', 2021)))
            if page == 'Commissions':
                self.assertEqual(app.text_input(key=input_key).value, '')
                self.assertEqual(app.text_input(key='date_paiement_AFRICAN LEASE_2021').value, '2021-02-01')
            else:
                self.assertIsNone(app.number_input(key=input_key).value)
            self.navigate(app, 'Comparaison')
            self.assertEqual(len(app.get('download_button')), 0)
        reload = self.navigate(self.app(), 'Suivi hebdomadaire')
        reload.sidebar.number_input(key='global_year').set_value(2021).run()
        self.assertIsNone(reload.number_input(key='weekly_obj_AFRICAN LEASE_2021_Janvier_1').value)
        self.assertEqual(reload.text_input(key='weekly_comment_AFRICAN LEASE_2021_Janvier_1').value, 'Conserver')
        reload.button(key='save_week_AFRICAN LEASE_2021_Janvier_1').click().run()
        self.assertIsNone(s.get_weekly_for_partner_mois('AFRICAN LEASE', 2021, 'Janvier')[1]['objectif'])
        self.assertEqual(s.get_data_for_partner_annee('AFRICAN LEASE', 2021)['ANNUEL']['realisation'], 700)
        self.navigate(reload, 'Production')
        self.assertIsNone(reload.number_input(key='monthly_obj_AFRICAN LEASE_2021_Janvier').value)
        self.assertIsNone(reload.number_input(key='annual_obj_AFRICAN LEASE_2021').value)
        reload.button(key='save_month_AFRICAN LEASE_2021_Janvier').click().run()
        reload.button(key='save_annual_AFRICAN LEASE_2021').click().run()
        self.assertIsNone(s.get_data_for_partner_annee('AFRICAN LEASE', 2021)['ANNUEL']['objectif'])
        self.navigate(reload, 'Commissions')
        self.assertEqual(reload.text_input(key='commission_reversee_AFRICAN LEASE_2021').value, '')
        reload.button(key='save_commissions_AFRICAN LEASE_2021').click().run()
        self.assertNotIn('commission_reversee', s.get_text_for_partner_annee('AFRICAN LEASE', 2021))
        from io import BytesIO
        import openpyxl
        from reporting import export_excel, export_presentation
        from pptx import Presentation
        workbook = openpyxl.load_workbook(BytesIO(export_excel(s, 'AFRICAN LEASE', [2021])), data_only=True)
        self.assertEqual(list(workbook['Historique'].values)[1][1:3], (None, 0))
        self.assertEqual(list(workbook['Hebdomadaire'].values)[1][3:6], (None, 0, 'Conserver'))
        deck = Presentation(BytesIO(export_presentation(s, 'AFRICAN LEASE', 2021)))
        all_text = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertNotIn('125 000', all_text)

# Empêche unittest de redécouvrir les 16 tests importés.
del _Base

if __name__ == '__main__':
    unittest.main()
