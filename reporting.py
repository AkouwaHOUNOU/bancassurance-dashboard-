"""Exports en mémoire, indépendants de Streamlit et sans écriture en production.

Les totaux utilisent le même service que les KPI. Les données absentes restent
vides dans Excel ; les anciennes valeurs annuelles figurent séparément.
"""
from io import BytesIO
import textwrap

import pandas as pd
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from database_service import MOIS, format_fcfa, format_pct

GOLD = 'D4A72C'
BLUE = '002B49'


def export_excel(service, partner: str, years: list[int]) -> bytes:
    """Conserve la feuille Historique et ses noms de colonnes historiques."""
    output = BytesIO()
    texts, weeks, originals = [], [], []
    for year in sorted(set(years)):
        for field, value in service.get_text_for_partner_annee(partner, year).items():
            texts.append({'Partenaire': partner, 'Année': year, 'Champ': field, 'Valeur': value})
        for month in MOIS:
            for week, data in service.get_weekly_for_partner_mois(partner, year, month).items():
                weeks.append({'Année': year, 'Mois': month, 'Semaine': week,
                              'Objectif': data['objectif'], 'Réalisation': data['realisation'],
                              'Commentaire': data['commentaire']})
        for kind, value in service.get_data_for_partner_annee(partner, year).get('ANNUEL', {}).items():
            originals.append({'Année': year, 'Type': kind, 'Valeur enregistrée': value})
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        service.get_comparison_dataframe(partner, sorted(set(years))).to_excel(writer, index=False, sheet_name='Historique')
        pd.DataFrame(texts, columns=['Partenaire', 'Année', 'Champ', 'Valeur']).to_excel(writer, index=False, sheet_name='Profil et commissions')
        pd.DataFrame(weeks, columns=['Année', 'Mois', 'Semaine', 'Objectif', 'Réalisation', 'Commentaire']).to_excel(writer, index=False, sheet_name='Hebdomadaire')
        pd.DataFrame(originals, columns=['Année', 'Type', 'Valeur enregistrée']).to_excel(writer, index=False, sheet_name='Totaux manuels conservés')
        for sheet in writer.book:
            sheet.freeze_panes = 'B2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color=GOLD)
                cell.fill = PatternFill('solid', fgColor=BLUE)
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = 25
            # Du texte utilisateur doit rester du texte, jamais une formule Excel.
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
    return output.getvalue()


def export_presentation(service, partner: str, year: int) -> bytes:
    """Rapport de l'exercice global : KPI cohérents, graphiques, profil complet."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    def slide(title, body=''):
        item = prs.slides.add_slide(prs.slide_layouts[6])
        item.background.fill.solid()
        item.background.fill.fore_color.rgb = RGBColor.from_string(BLUE)
        title_box = item.shapes.add_textbox(Inches(.5), Inches(.35), Inches(12.2), Inches(.8))
        title_box.text = title
        for paragraph in title_box.text_frame.paragraphs:
            paragraph.font.size = Pt(26)
            paragraph.font.color.rgb = RGBColor.from_string(GOLD)
        if body:
            box = item.shapes.add_textbox(Inches(.6), Inches(1.4), Inches(12), Inches(5.5))
            box.text = body
            box.text_frame.word_wrap = True
            for paragraph in box.text_frame.paragraphs:
                paragraph.font.size = Pt(18)
                paragraph.font.color.rgb = RGBColor(255, 255, 255)
        return item

    slide('Bancassurance NSIA', f'{partner} — Exercice {year}\nDonnées enregistrées en SQLite')
    kpi = service.get_gap_taux_historical(partner, year)
    slide(f'Tableau de bord — {year}', '\n'.join([
        f"Objectif : {format_fcfa(kpi['objectif_N'])}",
        f"Réalisation : {format_fcfa(kpi['realisation_N'])}",
        f"GAP (objectif − réalisation) : {format_fcfa(kpi['gap'])}",
        f"Taux de réalisation : {format_pct(kpi['taux'])}",
    ]))
    raw = service.get_data_for_partner_annee(partner, year)
    monthly = [raw.get(month, {}) for month in MOIS]
    # None garde les cellules absentes vides dans les graphiques natifs.
    obj = [entry.get('objectif') for entry in monthly]
    real = [entry.get('realisation') for entry in monthly]
    datasets = [
        ('Objectif et réalisation mensuels', [('Objectif', obj), ('Réalisation', real)], XL_CHART_TYPE.COLUMN_CLUSTERED),
        ('Taux de réalisation mensuel (%)', [('Taux (%)', [(r / o * 100) if o and r is not None else None for o, r in zip(obj, real)])], XL_CHART_TYPE.COLUMN_CLUSTERED),
    ]
    if any(value is not None and value > 0 for value in obj) and all(value is None or value >= 0 for value in obj):
        datasets.append(('Répartition mensuelle des objectifs', [('Objectif', obj)], XL_CHART_TYPE.PIE))
    for title, series, chart_type in datasets:
        item = slide(f'{title} — {year}')
        chart_data = CategoryChartData()
        chart_data.categories = MOIS
        for name, values in series:
            chart_data.add_series(name, values)
        chart = item.shapes.add_chart(chart_type, Inches(.6), Inches(1.4), Inches(12), Inches(5.4), chart_data).chart
        chart.has_legend = len(series) > 1
        for index, chart_series in enumerate(chart.series):
            chart_series.format.fill.solid()
            chart_series.format.fill.fore_color.rgb = RGBColor.from_string(GOLD if index == 0 else '47A1D8')
    # Pagination conservatrice : pas de texte tronqué pour les longs profils.
    text = service.get_text_for_partner_annee(partner, year)
    for field, value in text.items():
        if field == 'annual_mode' or not value:
            continue
        lines = textwrap.wrap(str(value), width=90, replace_whitespace=False) or ['']
        for offset in range(0, len(lines), 12):
            title = field.replace('_', ' ').capitalize()
            suffix = ' (suite)' if offset else ''
            slide(f'{title} — {year}{suffix}', '\n'.join(lines[offset:offset + 12]))
    output = BytesIO()
    prs.save(output)
    return output.getvalue()
